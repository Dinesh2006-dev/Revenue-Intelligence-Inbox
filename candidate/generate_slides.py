import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    os.makedirs('candidate/slides', exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(13, 17, 23)        # #0D1117
    SURFACE_COLOR = RGBColor(22, 27, 34)   # #161B22
    SURFACE2_COLOR = RGBColor(33, 38, 45)  # #21262D
    BORDER_COLOR = RGBColor(48, 54, 61)    # #30363D
    TEXT_COLOR = RGBColor(230, 237, 243)   # #E6EDF3
    TEXT_MUTED = RGBColor(139, 148, 158)   # #8B949E
    ACCENT_BLUE = RGBColor(88, 166, 255)   # #58A6FF
    ACCENT_GREEN = RGBColor(63, 185, 80)   # #3FB950
    ACCENT_AMBER = RGBColor(210, 153, 34)  # #D29922
    ACCENT_RED = RGBColor(248, 81, 73)     # #F85149

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, subtitle_text, slide_number):
        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title_text
        p1.font.name = "Arial"
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_COLOR

        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.name = "Arial"
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

        # Slide Number Badge
        num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.5), Inches(1.0), Inches(0.5))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.text = f"{slide_number} / 5"
        np.alignment = PP_ALIGN.RIGHT
        np.font.name = "Arial"
        np.font.size = Pt(11)
        np.font.bold = True
        np.font.color.rgb = ACCENT_BLUE

    # ── SLIDE 1: Problem, Objective & Product Overview ────────────────────────
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    add_header(slide1, "Revenue Intelligence Inbox: Product Architecture & Challenge",
               "SpendNest Sales Management Solution · Volopay Assessment", 1)

    cards_data_1 = [
        ("The Core Manager Challenge", ACCENT_BLUE, [
            "Sales managers are inundated with calls, calendar meetings, and disjointed CRM notes.",
            "Asynchronous arrival: Transcripts arrive late (e.g., meeting M07 arriving in revision 2).",
            "Noisy and deceptive metadata: Stored outcome labels (e.g., 'voicemail') frequently conflict with actual spoken dialog.",
            "Multi-objective conversations: Single calls contain both technical vendor evaluation and customer retention issues."
        ]),
        ("Product Mission & Solution", ACCENT_GREEN, [
            "Single Unified Workspace: Ingests 44 distinct interactions with zero duplicate inflation.",
            "Deterministic Rule Engine: Runs 100% locally with zero paid API dependencies or credentials.",
            "Transparent Evidence Anchoring: Every classification is cited to an exact transcript turn or verbatim quotation.",
            "Robust Contract Replay: Fully respects revision ordering, deduplication keys (source_type, source_id), and manual overrides."
        ]),
        ("Core Capabilities Delivered", ACCENT_AMBER, [
            "Executive Overview: Real-time telemetry, transcript coverage (40.9%), and connected-call rate (36.7%).",
            "Split-View Inbox: Fast filtering, metadata inspector, and interactive routing controls.",
            "Manager Daily Brief: Synthesized intelligence for 21 August 2026 across coaching and escalations.",
            "Audit & Export Hub: Complete machine-readable manifest (JSON/CSV) and verified file packages."
        ])
    ]

    for i, (title, color, bullets) in enumerate(cards_data_1):
        left = Inches(0.8 + i * 4.0)
        top = Inches(1.8)
        width = Inches(3.7)
        height = Inches(5.0)

        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE_COLOR
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = color

        for b in bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.name = "Arial"
            p_b.font.size = Pt(10.5)
            p_b.font.color.rgb = TEXT_COLOR
            p_b.space_before = Pt(8)

    # ── SLIDE 2: Dataset, Methodology & Metric Definitions ────────────────────
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Data Contract Rigor, Baseline Telemetry & Metric Definitions",
               "Ground Truth Analysis across 44 Applied Source Records", 2)

    cards_data_2 = [
        ("Dataset Structure & Coverage", ACCENT_BLUE, [
            "Total Ingested Records: Exactly 44 unique records (30 calls, 14 meetings).",
            "Initial Delivery (Rev 1): 17 usable conversation texts (10 calls, 7 meetings).",
            "Update Delivery (Rev 2): M07 delivers 27 turns, expanding usable coverage to 18 (40.9%).",
            "Unprocessed / Silent Records: 26 records (20 background calls, 6 silent/cancelled meetings) retained as metadata without hallucinated analysis."
        ]),
        ("Authoritative Metric Formulations", ACCENT_GREEN, [
            "Transcript Coverage: 18 / 44 = 40.9% (Usable texts divided by selected records).",
            "Source-Reported Connected Rate: 11 / 30 = 36.7% (Calls where is_connected=true divided by non-null is_connected; unknown count = 0).",
            "Separation of Concepts: Upstream stored connection status is explicitly separated from qualitative transcript analysis.",
            "Null Value Handling: Null is strictly preserved as 'value absent', never coerced into zero or false."
        ]),
        ("Deduplication & Replay Contract", ACCENT_AMBER, [
            "Logical Identity: Strictly the composite key (source_type, source_id).",
            "Anti-Pattern Prevention: Never deduplicate on title, company name, phone, or array position.",
            "Revision Precedence: Higher revision replaces full payload; older revision is ignored.",
            "Override Invariance: Manual routing decisions are preserved across subsequent re-imports and application restarts."
        ])
    ]

    for i, (title, color, bullets) in enumerate(cards_data_2):
        left = Inches(0.8 + i * 4.0)
        top = Inches(1.8)
        width = Inches(3.7)
        height = Inches(5.0)

        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE_COLOR
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = color

        for b in bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.name = "Arial"
            p_b.font.size = Pt(10.5)
            p_b.font.color.rgb = TEXT_COLOR
            p_b.space_before = Pt(8)

    # ── SLIDE 3: Three Important Findings with Exact Evidence ─────────────────
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Core Conversational Discoveries & Verifiable Citations",
               "Grounded Source Evidence Directly Influencing Interface and Routing Rules", 3)

    findings = [
        ("Finding 1: LedgerBridge (M02) is a Vendor Evaluation, Not a Sales Pitch", ACCENT_GREEN,
         "Turn 2 & Turn 31",
         "Supplier A: 'Then I will cover the market strategy as well as the technical integration... The basic setup begins at EUR 2,600; a more supported implementation starts at EUR 6,500...'",
         "Impact: Precluded false pipeline inflation. Routed to 'Vendor / Supplier' rather than Sales Opportunity, preserving quota integrity."),

        ("Finding 2: Willowfield Trust (M07) Revealed 4 Severe Software Defects", ACCENT_RED,
         "Turns 3, 8, 19, 21",
         "Customer A: 'When I use Mark as paid, the first payment shows INR 18,200... The amount from the previous item appears in the next one.' (Also: description edit fail, document upload fail, stuck reimbursements).",
         "Impact: Triggered immediate Customer Escalation alert, engineering bug ticket creation, and stale analysis banner."),

        ("Finding 3: Internal Ops Training (M10) Identified via Speaker Topology", ACCENT_BLUE,
         "Turns 4 & 5",
         "Team member B: 'I have access to a demo account... but not back-office super-admin access.' Team member A: 'Operations uses it to create accounts, control enabled features and manage billing.'",
         "Impact: Speaker detection identified all participants as internal staff; routed to 'Internal Training / Ops', keeping sales inbox uncluttered.")
    ]

    for i, (title, color, cite, quote, impact) in enumerate(findings):
        left = Inches(0.8 + i * 4.0)
        top = Inches(1.8)
        width = Inches(3.7)
        height = Inches(5.0)

        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE_COLOR
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p_cite = tf.add_paragraph()
        p_cite.text = f"Source Evidence: {cite}"
        p_cite.font.name = "Arial"
        p_cite.font.size = Pt(10)
        p_cite.font.bold = True
        p_cite.font.color.rgb = ACCENT_AMBER
        p_cite.space_before = Pt(6)

        p_q = tf.add_paragraph()
        p_q.text = f'"{quote}"'
        p_q.font.name = "Arial"
        p_q.font.size = Pt(9.5)
        p_q.font.italic = True
        p_q.font.color.rgb = TEXT_MUTED
        p_q.space_before = Pt(6)

        p_imp = tf.add_paragraph()
        p_imp.text = impact
        p_imp.font.name = "Arial"
        p_imp.font.size = Pt(10)
        p_imp.font.color.rgb = TEXT_COLOR
        p_imp.space_before = Pt(8)

    # ── SLIDE 4: Difficult Routing Decision & Rejected Alternative ─────────────
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Difficult Routing Case & Deliberately Rejected Decision",
               "Architectural Trade-offs, Edge-Case Governance, and Honest Uncertainty", 4)

    cards_data_4 = [
        ("Difficult Case: M04 (Cedarwave Distribution)", ACCENT_AMBER, [
            "Multi-Faceted Tension: Cedarwave is an existing customer suffering physical prepaid card rejection in stores.",
            "Conflicting Intents: Customer demands proper credit cards for their INR 1.3 crore bank facility, while rep attempts to pitch a prepaid migration.",
            "Organizational Friction: Customer also reports broken user department hierarchy and admin portal confusion.",
            "Resolution: Avoided categorizing as a simple 'Sales Migration'. Classified as 'Customer Concern / Retention Risk' with an action item to investigate bank credit feed feasibility."
        ]),
        ("Deliberately Rejected Decision: Synthetic Pipeline Valuation", ACCENT_RED, [
            "Alternative Considered: Using regex/LLM to extract currency figures mentioned in conversations into an aggregate CRM pipeline total.",
            "Why Rejected: Currency mentions in conversation represent customer operational costs (e.g., M03 INR 10-13L travel pool, M04 INR 1.3 crore credit limit, M06 INR 1.3L invoice), NOT SpendNest contract revenue!",
            "Unacceptable Risk: Summing these values would manufacture a deceptive multi-crore pipeline metric that misleads leadership.",
            "Accepted Trade-off: Displayed qualitative deal signals and confirmed next steps without fabricated financial totals."
        ]),
        ("Rule-Based Determinism vs Live LLM Calls", ACCENT_BLUE, [
            "Zero Cloud Cost: Runs entirely client-side without API keys, token limits, or network latency.",
            "Auditable Reasoning: Every routing suggestion is backed by transparent deterministic checks.",
            "Offline Operationality: Complies with strict data governance; sensitive conversation text never leaves the browser.",
            "Human-in-the-Loop Override: Managers retain 100% control to adjust routes with immutable audit logs."
        ])
    ]

    for i, (title, color, bullets) in enumerate(cards_data_4):
        left = Inches(0.8 + i * 4.0)
        top = Inches(1.8)
        width = Inches(3.7)
        height = Inches(5.0)

        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE_COLOR
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        for b in bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.name = "Arial"
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = TEXT_COLOR
            p_b.space_before = Pt(6)

    # ── SLIDE 5: Replay Testing, Product Decisions & Next Steps ────────────────
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Automated Replay Verification, Known Limitations & Roadmap",
               "100% Contract Test Suite Execution and Future Evolution", 5)

    cards_data_5 = [
        ("Automated Contract Verification (6/6 Pass)", ACCENT_GREEN, [
            "Test 1 (Initial Import): Ingested 44 unique records, 17 usable texts, 0 duplicates.",
            "Test 2 (Reversed Order Re-import): Idempotency proven; record count remained exactly 44.",
            "Test 3 (Manual Override): Changed M07 to 'customer_escalation' with audit note.",
            "Test 4 (Revision 2 Update): Ingested M07 update (27 turns); manual route strictly preserved.",
            "Test 5 (Older Rev 1 Re-import): Revision precedence proven; Rev 1 did not overwrite Rev 2.",
            "Test 6 (Browser Reload): State and overrides fully restored from LocalStorage."
        ]),
        ("Known Limitations & Disclosed Boundaries", ACCENT_AMBER, [
            "Temporal Extract Scope: All activity represents a single local date (21 August 2026); trendlines across weeks cannot be derived.",
            "Background Call Visibility: 20 calls lack transcripts due to upstream delivery redaction; classified strictly on metadata.",
            "No Authoritative Staff Directory: Rep roles are inferred from call direction and meeting titles.",
            "No CRM Stage Join: Conversation signals are presented without unverified CRM opportunity IDs."
        ]),
        ("Future Roadmap & Enhancements", ACCENT_BLUE, [
            "Bi-Directional CRM Sync: Webhook connector to push manual overrides back to Salesforce/HubSpot.",
            "Local Small-Language Model: Optional on-device WebGPU model (e.g., Gemma 2B) for zero-latency semantic search.",
            "Audio Waveform Viewer: Integration with audio player if un-redacted media is provided in future batches.",
            "Automated Engineering Webhook: Push verified customer bug passages directly to Jira/Linear."
        ])
    ]

    for i, (title, color, bullets) in enumerate(cards_data_5):
        left = Inches(0.8 + i * 4.0)
        top = Inches(1.8)
        width = Inches(3.7)
        height = Inches(5.0)

        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE_COLOR
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        for b in bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.name = "Arial"
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = TEXT_COLOR
            p_b.space_before = Pt(6)

    # Save presentation
    pptx_path = 'candidate/slides/revenue_intelligence_inbox_slides.pptx'
    prs.save(pptx_path)
    print(f"Generated 5-slide PPTX at {pptx_path}")

if __name__ == '__main__':
    create_presentation()
